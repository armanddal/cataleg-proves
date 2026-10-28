#!/usr/bin/env python3
"""Convertidor genèric «pàgina d'informació del catàleg -> .pptx A4 editable».

Ús:
  python3 scripts/pptx_pagina.py --grup 8                 # totes les pàgines d'informació del grup 8
  python3 scripts/pptx_pagina.py --grup 8 --pagina 4      # només la pàgina renderitzada 4
  python3 scripts/pptx_pagina.py --grup 8 --previsualitza # a més: PNG de previsualització (soffice + pdftoppm)

Com funciona: carrega sortida/gN.html al Chrome headless (amb el JS d'ajust del generador, de manera que el mode
«fluid» ja està aplicat) i hi injecta un extractor que, per cada `.page.info`, retorna la geometria i l'estil dels
elements (capçalera, columnes de text, fotos, referències). Amb això escriu un .pptx per pàgina, amb cada element com a
objecte separat i editable. No toca sortida/gN.* (només el llegeix).

El peu (data / logo / número / enllaç) NO s'inclou: el generador el posa al damunt (franja de baix buida).
Pensat per a Google Slides: formes i text simples, sense efectes ni SmartArt.
Fonts: títols 'Roboto Condensed' (Google Slides la té); text 'Arial' amb la mida reduïda perquè Arial és més ampla que
l'Arial Narrow del HTML (el factor es tria per pàgina simulant el tall de línies amb les mètriques reals d'Arial).

CALIBRATGE AMB GOOGLE SLIDES (v2, 28/10/2026)
Origen de les dades: els PDF exportats de Slides de les 20 diapositives de G8 (treball/slides/<id>.pdf, ja convertides des
dels .pptx v1) comparats amb la pàgina objectiu de sortida/g8.pdf (pdftotext -bbox-layout, caixes de línia i de paraula).
  1. INTERLINEAT (el defecte gros). Google Slides NO respecta `a:lnSpc/a:spcPts` tal qual: el pas entre línies de continuació
     d'un paràgraf surt EXACTAMENT 1,200 x el valor demanat (mesurat sobre 14 combinacions mida/interlineat diferents,
     p05, p10, p15, p17, p21, p22...; mediana = 1,200 en totes; independent de la mida de lletra i de si hi ha vinyeta).
     `a:spcBef` sí que es respecta (pas entre paràgrafs = 1,2 x L + spcBef). Per això escrivim L / SLIDES_LNSPC_K amb
     K = 1,20. Sense això cada columna s'allarga ~20 % en vertical (p05 +64 pt, p16 queda tallada per la vora de la pàgina).
     LibreOffice sí que respecta spcPts, per això la v1 hi semblava bé: amb K=1,2 el render de LibreOffice queda un 17 %
     més atapeït que l'objectiu; només Slides és la referència bona.
  2. AMPLE DEL TEXT. Les paraules d'Slides surten amb el mateix ample que a l'objectiu quan escala = 0,82 (Arial Narrow ~
     0,82 x Arial: mesurat com a mediana de l'ample de paraula Slides/objectiu = escala/0,82, p.ex. 1,097 amb 0,90). Es
     manté la tria «la més gran que no afegeix línies», però es limita a SCALE_MAX = 0,85 (abans 0,90; compromís: 0,82 iguala l ample però encongeix l alçada de glif) per no passar-se.
     Slides no encongeix el text (noAutofit) i els insets ja són 0 explícits (confirmat: l'x del text coincideix a <0,5 pt).
  2b. MARGE DE TALL. L'estimació amb PIL no deixava cap marge i una línia de 589,98 px dins 590,49 es trencava a Slides i a
     LibreOffice (p08, +2 línies): l'ample útil es redueix un 0,5 % (FIT_MARGIN).
  3. SALVAGUARDA D'ALÇADA. Per cada columna s'estima l'alçada (línies segons les mètriques d'Arial x interlineat objectiu +
     espais entre paràgrafs); si no cap fins a la foto següent / el peu, es redueix la lletra (mín. 6 pt) fins que càpiga.
  4. CERCLES AMB NÚMERO. L'el·lipse reserva per al text només el quadrat inscrit (0,707 x el diàmetre) i Slides trenca
     «10»/«11» en dues línies encara que sigui wrap=none. El número ara és un quadre de text a part (ample = 2 x diàmetre,
     centrat) agrupat amb el cercle.
  5. Les referències ①-⑳ no són a l'Arial: Slides usa MS PGothic de reserva, que fa la línia ~1 pt més alta (no calibrable).
"""
import argparse, base64, html as htmlmod, io, json, os, re, shutil, subprocess, sys, tempfile
from urllib.parse import unquote, urlparse

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from renderitza import CHROME_ARGS  # noqa: E402

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Cm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

PX = 37.795275591          # px CSS per cm
TIT_FONT, TXT_FONT = 'Roboto Condensed', 'Arial'
SLIDES_LNSPC_K = 1.20      # Slides pinta l'interlineat exacte (spcPts) x1,20: s'escriu L/K (veure CALIBRATGE)
SCALE_MAX = 0.85           # mida màxima del text respecte a l'HTML (Arial Narrow ~0,82 x Arial)
FIT_MARGIN = 0.995         # marge d'ample en l'estimació de línies: Slides/LibreOffice trenquen on PIL ho dona per just (p08)
MIN_PT = 6.0               # lletra mínima de la salvaguarda d'alçada
FONT_ARIAL = {False: '/System/Library/Fonts/Supplemental/Arial.ttf', True: '/System/Library/Fonts/Supplemental/Arial Bold.ttf'}
FONT_NARROW_B = '/System/Library/Fonts/Supplemental/Arial Narrow Bold.ttf'

# ------------------------------------------------------------------ extractor (s'executa dins el Chrome)
EXTRACT_JS = r'''
<script>
window.addEventListener('load', function(){ setTimeout(function(){
  var all = Array.from(document.querySelectorAll('.page'));
  var out = [];
  function rgb(c){ var m = c.match(/[\d.]+/g); if (!m) return null; var a = m.length > 3 ? parseFloat(m[3]) : 1; return a === 0 ? null : [Math.round(m[0]), Math.round(m[1]), Math.round(m[2])]; }
  all.forEach(function(pg, ix){
    if (!pg.classList.contains('info')) return;
    var pr = pg.getBoundingClientRect();
    function R(e){ var b = e.getBoundingClientRect(); return {x: b.left - pr.left, y: b.top - pr.top, w: b.width, h: b.height}; }
    function inPeu(e){ return !!e.closest('.peu'); }
    var P = {peu: null, n: ix + 1, classes: pg.className, why: pg.dataset.why || '', cap: null, chip: null, tes: null, tx: [], lts: [], imgs: [], bigs: [], pars: [], bloques: [], text: ''};
    var cap = pg.querySelector('.cap');
    if (cap) P.cap = Object.assign(R(cap), {bg: rgb(getComputedStyle(cap).backgroundColor)});
    var peu = pg.querySelector('.peu');
    if (peu) P.peu = R(peu);
    var chip = pg.querySelector('.chip');
    if (chip) { var cs = getComputedStyle(chip); P.chip = Object.assign(R(chip), {bg: rgb(cs.backgroundColor), text: chip.innerText.trim(), fs: parseFloat(cs.fontSize), color: rgb(cs.color), pt: parseFloat(cs.paddingTop)}); }
    var tes = pg.querySelector('.cap-tit .t-es');
    if (tes) { var ts = getComputedStyle(tes); P.tes = Object.assign(R(tes), {text: tes.textContent.trim(), fs: parseFloat(ts.fontSize), color: rgb(ts.color), bb: parseFloat(ts.borderBottomWidth), bbc: rgb(ts.borderBottomColor)}); }
    pg.querySelectorAll('.cap-tit .t-x').forEach(function(t){
      var lt = t.querySelector('.lt'), ts = getComputedStyle(t), s = '';
      t.childNodes.forEach(function(c){ if (c !== lt) s += c.textContent; });
      var o = Object.assign(R(t), {text: s.trim(), fs: parseFloat(ts.fontSize), color: rgb(ts.color), lt: null});
      if (lt) { var ls = getComputedStyle(lt); o.lt = Object.assign(R(lt), {text: lt.textContent.trim(), fs: parseFloat(ls.fontSize), color: rgb(ls.borderTopColor) || rgb(ls.color)}); }
      P.tx.push(o);
    });
    // lletres E/F/D de les columnes de text
    pg.querySelectorAll('.col-lang .lt, .bl-col .lt').forEach(function(lt){
      var ls = getComputedStyle(lt);
      P.lts.push(Object.assign(R(lt), {text: lt.textContent.trim(), fs: parseFloat(ls.fontSize), color: rgb(ls.borderTopColor) || rgb(ls.color)}));
    });
    // imatges
    pg.querySelectorAll('img').forEach(function(im){
      if (inPeu(im)) return;
      var b = R(im), cs = getComputedStyle(im);
      if (b.w < 2 || b.h < 2) return;
      P.imgs.push(Object.assign(b, {src: im.currentSrc || im.src, nw: im.naturalWidth, nh: im.naturalHeight, fit: cs.objectFit, pos: cs.objectPosition}));
    });
    // referències grans (cercle damunt de la foto)
    pg.querySelectorAll('i.ref.big, i.bl-n').forEach(function(r){
      var cs = getComputedStyle(r);
      P.bigs.push(Object.assign(R(r), {text: r.textContent.trim(), fs: parseFloat(cs.fontSize), color: rgb(cs.color), line: rgb(cs.borderTopColor), bg: rgb(cs.backgroundColor), round: cs.borderRadius !== '0px', lw: parseFloat(cs.borderTopWidth)}));
    });
    // paràgrafs de text
    var cols = Array.from(pg.querySelectorAll('.info-col, .bl-col, .bl-text'));
    pg.querySelectorAll('.info-col h3, .info-col li, .bl-col h4, .bl-col li, .bl-text p').forEach(function(e){
      var cs = getComputedStyle(e), runs = [];
      (function walk(n){ n.childNodes.forEach(function(c){
        if (c.nodeType === 3) runs.push({t: c.textContent});
        else if (c.nodeType === 1) { if (c.classList.contains('ref')) runs.push({ref: c.textContent.trim()}); else if (c.classList.contains('lt')) {} else walk(c); }
      }); })(e);
      var lh = parseFloat(cs.lineHeight); if (isNaN(lh)) lh = parseFloat(cs.fontSize) * 1.2;
      var b = R(e), col = e.closest('.info-col, .bl-col, .bl-text');
      var kind = e.tagName.toLowerCase(), list = e.parentNode && e.parentNode.tagName ? e.parentNode.tagName.toLowerCase() : '';
      var tx = (e.textContent || '').trim();
      var frs = Array.from(e.getClientRects()).map(function(q){ return {x: q.left - pr.left, y: q.top - pr.top, w: q.width, h: q.height}; });
      if (frs.length > 1 && frs.some(function(f, i){ return i && Math.abs(f.x - frs[0].x) > 20; })) {      // bloc partit entre columnes (CSS columns): un paràgraf per fragment
        var parts = frs.map(function(){ return []; });
        function frag(rr){ var cx = rr.left - pr.left + rr.width / 2, cy = rr.top - pr.top + rr.height / 2, best = 0, bd = 1e9;
          frs.forEach(function(f, i){ var dx = Math.max(f.x - cx, 0, cx - f.x - f.w), dy = Math.max(f.y - cy, 0, cy - f.y - f.h), d = dx * dx + dy * dy; if (d < bd) { bd = d; best = i; } }); return best; }
        (function walk(n){ n.childNodes.forEach(function(c){
          if (c.nodeType === 3) { var re = /\S+\s*/g, m, tx = c.textContent;
            while ((m = re.exec(tx))) { var rg = document.createRange(); rg.setStart(c, m.index); rg.setEnd(c, m.index + m[0].replace(/\s+$/, '').length || m[0].length); var q = rg.getClientRects()[0]; if (!q) continue; parts[frag(q)].push({t: m[0]}); } }
          else if (c.nodeType === 1) { if (c.classList.contains('ref')) parts[frag(c.getBoundingClientRect())].push({ref: c.textContent.trim()}); else if (!c.classList.contains('lt')) walk(c); }
        }); })(e);
        frs.forEach(function(f, i){ if (!parts[i].length) return;
          P.pars.push({kind: kind, cont: i > 0, list: list, runs: parts[i], x: f.x, y: f.y, w: f.w, h: f.h, fs: parseFloat(cs.fontSize), lh: lh, color: rgb(cs.color),
            bold: parseInt(cs.fontWeight) >= 600, align: cs.textAlign, pl: parseFloat(cs.paddingLeft), upper: cs.textTransform === 'uppercase',
            col: cols.indexOf(col), lines: Math.round(f.h / lh), vermell: rgb(getComputedStyle(pg.querySelector('i.ref') || e).borderTopColor)}); });
        return;
      }
      P.pars.push({kind: kind, list: list, runs: runs, x: b.x, y: b.y, w: b.w, h: b.h, fs: parseFloat(cs.fontSize), lh: lh, color: rgb(cs.color),
                   bold: parseInt(cs.fontWeight) >= 600, align: cs.textAlign, pl: parseFloat(cs.paddingLeft), upper: cs.textTransform === 'uppercase',
                   col: cols.indexOf(col), lines: Math.round(b.h / lh), vermell: rgb(getComputedStyle(pg.querySelector('i.ref') || e).borderTopColor)});
    });
    // rectangles amb vora (mode blocs)
    pg.querySelectorAll('.bl-foto, .bl-bloc').forEach(function(e){
      var cs = getComputedStyle(e); P.bloques.push(Object.assign(R(e), {line: rgb(cs.borderTopColor), lw: parseFloat(cs.borderTopWidth)}));
    });
    // text visible de la pàgina (sense capçalera ni peu) per comprovar que no es perd res
    var clone = pg.cloneNode(true); clone.querySelectorAll('.peu,.cap').forEach(function(x){ x.remove(); });
    P.text = clone.textContent.replace(/\s+/g, ' ').trim();
    out.push(P);
  });
  var s = btoa(unescape(encodeURIComponent(JSON.stringify(out))));
  var pre = document.createElement('pre'); pre.id = 'pptx_json'; pre.textContent = s; document.body.appendChild(pre);
}, 50); });
</script>
'''


def extreu(html_path):
    """Executa el HTML al Chrome (amb l'ajust del generador) i retorna la llista de pàgines d'informació extretes."""
    doc = open(html_path, encoding='utf8').read()
    base = '<base href="file://%s/">' % os.path.dirname(os.path.abspath(html_path))
    doc = doc.replace('<head>', '<head>' + base, 1).replace('</body>', EXTRACT_JS + '</body>')
    tmpd = tempfile.mkdtemp(prefix='pptx-ext-')
    perfil = tempfile.mkdtemp(prefix='chrome-pptx-')
    try:
        f = os.path.join(tmpd, 'x.html'); open(f, 'w', encoding='utf8').write(doc)
        r = subprocess.run(CHROME_ARGS() + [f'--user-data-dir={perfil}', '--window-size=794,1123', '--virtual-time-budget=60000',
                                            '--dump-dom', 'file://' + f], capture_output=True, text=True, timeout=300)
    finally:
        shutil.rmtree(perfil, ignore_errors=True); shutil.rmtree(tmpd, ignore_errors=True)
    m = re.search(r'<pre id="pptx_json">(.*?)</pre>', r.stdout, re.S)
    if not m: raise RuntimeError('El Chrome no ha retornat l\'extracció: ' + r.stderr[-400:])
    return json.loads(base64.b64decode(htmlmod.unescape(m.group(1))).decode('utf8'))


# ------------------------------------------------------------------ utilitats pptx
def cm(px): return px / PX
def rgbc(c, defecte=(0, 0, 0)): return RGBColor(*(c or defecte))
def pt(px): return px * 0.75


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.6, nom=None, shape=MSO_SHAPE.RECTANGLE):
    s = sl.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    st = s._element.find(qn('p:style'))
    if st is not None: s._element.remove(st)
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    if nom: s.name = nom
    return s


def set_run(r, size, color, bold=False, font=TXT_FONT):
    r.font.size = Pt(round(size * 100) / 100); r.font.bold = bold; r.font.name = font; r.font.color.rgb = color


def tbox(sl, x, y, w, h, nom, anchor=MSO_ANCHOR.TOP, wrap=True):
    t = sl.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)); t.name = nom
    tf = t.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    bp = tf._txBody.find(qn('a:bodyPr'))
    for c in list(bp): bp.remove(c)
    etree.SubElement(bp, qn('a:noAutofit'))
    return t


def etiqueta(sl, x, y, w, h, text, size, color, nom, bold=False, font=TXT_FONT, fill=None, line=None, lw=0.75,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=MSO_SHAPE.RECTANGLE):
    s = rect(sl, x, y, w, h, fill, line, lw, nom, shape)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; set_run(r, size, color, bold, font)
    return s


def ordena_ppr(p):
    pPr = p._p.find(qn('a:pPr'))
    if pPr is None: return
    ordre = ['lnSpc', 'spcBef', 'spcAft', 'buClrTx', 'buClr', 'buSzTx', 'buSzPct', 'buSzPts', 'buFontTx', 'buFont',
             'buNone', 'buAutoNum', 'buChar', 'buBlip', 'tabLst', 'defRPr']
    kids = sorted(list(pPr), key=lambda e: ordre.index(etree.QName(e).localname))
    for k in kids: pPr.remove(k)
    for k in kids: pPr.append(k)


_img_cache = {}
def prepara_imatge(path, w_cm, h_cm, tmpdir):
    """Retorna (fitxer jpg/png reduït, (ample px, alt px)). Aplana la transparència sobre blanc; ≤1500 px i ≤~250 ppp."""
    clau = (path, round(w_cm, 1), round(h_cm, 1))
    if clau in _img_cache: return _img_cache[clau]
    im = Image.open(path)
    if im.mode in ('RGBA', 'LA') or (im.mode == 'P' and 'transparency' in im.info):
        im = im.convert('RGBA'); fons = Image.new('RGB', im.size, (255, 255, 255)); fons.paste(im, mask=im.split()[3]); im = fons
    else: im = im.convert('RGB')
    maxpx = min(1500, max(300, int(max(w_cm, h_cm) / 2.54 * 250)))
    if max(im.size) > maxpx:
        k = maxpx / max(im.size); im = im.resize((round(im.width * k), round(im.height * k)), Image.LANCZOS)
    out = os.path.join(tmpdir, 'i%d.jpg' % len(_img_cache))
    im.save(out, 'JPEG', quality=86)
    _img_cache[clau] = (out, im.size)
    return _img_cache[clau]


def ruta_local(src):
    return unquote(urlparse(src).path)


# ------------------------------------------------------------------ text: tall de línies simulat (tria del factor de lletra)
_fonts = {}
def amplada(text, size_px, bold):
    k = (bold, round(size_px * 4))
    if k not in _fonts: _fonts[k] = ImageFont.truetype(FONT_ARIAL[bold], max(1, round(size_px * 4)))
    return _fonts[k].getlength(text) / 4.0


def linies_arial(par, escala):
    """Nombre de línies que faria el paràgraf amb Arial a mida fs*escala dins l'ample disponible."""
    fs = par['fs'] * escala; disp = (par['w'] - par['pl']) * FIT_MARGIN - 1.0
    if par['list'] == 'ol': disp = (par['w'] - par['pl']) * FIT_MARGIN
    paraules = []
    for r in par['runs']:
        if 'ref' in r: paraules.append(('①', True))      # cercle
        else: paraules += [(w, False) for w in re.split(r'(?<=\s)', re.sub(r'\s+', ' ', r['t'])) if w]
    n, ample = 1, 0.0
    for w, es_ref in paraules:
        a = amplada(w.rstrip(), fs, par['bold']) if not es_ref else fs * 1.0
        asp = amplada(' ', fs, par['bold'])
        if ample + a > disp + 0.01 and ample > 0:
            n += 1; ample = 0.0
        ample += a + (asp if w.endswith(' ') else 0) + (fs * 0.2 if es_ref else 0)
    return n


def tria_escala(pars):
    for esc in [round(SCALE_MAX - 0.005 * i, 3) for i in range(0, 31)]:      # 0,85 ... 0,70 de mig en mig punt
        if all(linies_arial(p, esc) <= p['lines'] for p in pars): return esc
    return 0.72


def alt_estimada(ps, escala):
    """Alçada (px CSS) que ocuparà el grup de paràgrafs a Slides: línies d'Arial a aquesta escala x interlineat de l'HTML
    + els espais entre paràgrafs de l'HTML. Només és fiable si l'interlineat ja està corregit (SLIDES_LNSPC_K)."""
    h = 0.0; ant = None
    for p in ps:
        h += linies_arial(p, escala) * p['lh']
        if ant is not None: h += max(0.0, p['y'] - ant)
        ant = p['y'] + p['h']
    return h


def limit_inferior(P, x, w, y):
    """y màxima (px) fins on pot arribar un text que comença a (x, y) amb ample w: la foto següent que se li solapa
    en horitzontal, o el peu de la pàgina."""
    lim = P['peu']['y'] - 2 if P.get('peu') else 1123 - 60
    for im in P['imgs']:
        if im['y'] > y + 8 and im['x'] < x + w and im['x'] + im['w'] > x and im['y'] - 2 < lim: lim = im['y'] - 2
    for b in P['bigs']:
        if b['y'] > y + 8 and b['x'] < x + w and b['x'] + b['w'] > x and b['y'] - 2 < lim: lim = b['y'] - 2
    return lim


def escala_grup(P, ps, escala, x, w, y):
    """Salvaguarda: baixa la lletra (mín. MIN_PT) fins que la columna càpiga entre y i la foto següent / el peu."""
    lim = limit_inferior(P, x, w, y)
    min_fs = min(p['fs'] for p in ps)
    e_min = min(escala, MIN_PT / pt(min_fs))
    e = escala
    while alt_estimada(ps, e) > lim - y and e > e_min + 1e-6:
        e = max(e_min, e - 0.01)
    return e, alt_estimada(ps, e) > lim - y + 0.5


# ------------------------------------------------------------------ construcció de la diapositiva
CERCLES = {str(i + 1): chr(0x2460 + i) for i in range(20)}


def omple_pars(tb, pars, escala):
    tf = tb.text_frame
    primer = True; ant_bottom = None
    for p in pars:
        par = tf.paragraphs[0] if primer else tf.add_paragraph()
        pPr = par._p.get_or_add_pPr()
        bullet = p['kind'] == 'li' and not p.get('cont')
        if bullet:
            marL = max(0.0, cm(p['pl']))
            if p['list'] == 'ol':
                pPr.set('marL', str(Cm(marL))); pPr.set('indent', str(-Cm(marL)))
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')
            else:
                sang = min(marL, 0.08) if marL < 0.2 else 0.08
                pPr.set('marL', str(Cm(marL))); pPr.set('indent', str(-Cm(max(0.0, marL - sang))))
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
        else:
            if p['kind'] == 'li': pPr.set('marL', str(Cm(max(0.0, cm(p['pl']))))); pPr.set('indent', '0')
            etree.SubElement(pPr, qn('a:buNone'))
        par.line_spacing = Pt(round(pt(p['lh']) / SLIDES_LNSPC_K * 100) / 100)
        par.space_before = Pt(0 if ant_bottom is None else max(0.0, round(pt(p['y'] - ant_bottom) * 10) / 10))
        par.space_after = Pt(0)
        par.alignment = {'justify': PP_ALIGN.JUSTIFY, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(p['align'], PP_ALIGN.LEFT)
        ant_bottom = p['y'] + p['h']
        primer = False
        mida = pt(p['fs']) * escala
        runs = p['runs']
        for k, r in enumerate(runs):
            if 'ref' in r:
                run = par.add_run(); run.text = ' ' + CERCLES.get(r['ref'], '(%s)' % r['ref'])
                set_run(run, mida * 1.05, rgbc(p.get('vermell') or [204, 0, 0]), True, TXT_FONT)
            else:
                t = re.sub(r'\s+', ' ', r['t'])
                if k == 0: t = t.lstrip()
                if k == len(runs) - 1 or all('ref' not in x for x in runs[k + 1:]) and k == len(runs) - 1: t = t.rstrip()
                if not t: continue
                run = par.add_run(); run.text = t.upper() if p['upper'] else t
                set_run(run, mida, rgbc(p['color']), p['bold'], TXT_FONT)
        # elimina l'espai final del darrer run de text si va seguit només de referències buides
        ordena_ppr(par)


def construeix(P, sortida, tmpdir):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(21), Cm(29.7)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    avisos = []

    # --- capçalera
    if P['cap']:
        c = P['cap']; rect(sl, cm(c['x']), cm(c['y']), cm(c['w']), cm(c['h']), rgbc(c['bg']), nom='Banda vermella')
    if P['chip']:
        c = P['chip']
        s = etiqueta(sl, cm(c['x']), cm(c['y']), cm(c['w']), cm(c['h']), '', pt(c['fs']), rgbc(c['color'], (255, 255, 255)), 'Xip serie',
                     bold=True, fill=rgbc(c['bg']), anchor=MSO_ANCHOR.TOP)
        tf = s.text_frame; tf.word_wrap = True; tf.margin_top = Cm(cm(c['pt'] or 4)); tf.margin_left = tf.margin_right = Cm(0.05)
        p = tf.paragraphs[0]
        for r in list(p._p.findall(qn('a:r'))): p._p.remove(r)
        lines = [l.strip() for l in c['text'].split('\n') if l.strip()] or [c['text']]
        for i, l in enumerate(lines):
            par = p if i == 0 else tf.add_paragraph(); par.alignment = PP_ALIGN.CENTER
            r = par.add_run(); r.text = l; set_run(r, pt(c['fs']), rgbc(c['color'], (255, 255, 255)), True, TXT_FONT)
    if P['tes']:
        t = P['tes']; ample_cm = cm(t['w'])
        fs = pt(t['fs'])
        # Roboto Condensed Bold ≈ Arial Narrow Bold +4 %: es redueix la mida si el títol no hi cap en una línia
        f = ImageFont.truetype(FONT_NARROW_B, 200)
        while f.getlength(t['text'].upper() if False else t['text']) / 200 * fs * 1.04 / 72 * 2.54 > ample_cm - 0.2 and fs > 9: fs -= 0.25
        bb = t['bb'] or 0
        tb = tbox(sl, cm(t['x']), cm(t['y']), ample_cm, cm(t['h'] - bb), 'Titol ES', MSO_ANCHOR.MIDDLE)
        r = tb.text_frame.paragraphs[0].add_run(); r.text = t['text']; set_run(r, fs, rgbc(t['color'], (255, 255, 255)), True, TIT_FONT)
        if bb:
            y = cm(t['y'] + t['h'] - bb / 2)
            ln = sl.shapes.add_connector(1, Cm(cm(t['x'])), Cm(y), Cm(cm(t['x'] + t['w'])), Cm(y))
            st = ln._element.find(qn('p:style'))
            if st is not None: ln._element.remove(st)
            ln.line.color.rgb = rgbc(t['bbc'], (255, 255, 255)); ln.line.width = Pt(0.75); ln.name = 'Linia blanca'
    for t in P['tx']:
        x0 = t['x']
        if t['lt']:
            lt = t['lt']
            etiqueta(sl, cm(lt['x']), cm(lt['y']), cm(lt['w']), cm(lt['h']), lt['text'], pt(lt['fs']) * 0.9, rgbc(lt['color'], (255, 255, 255)),
                     'Quadret ' + lt['text'], bold=True, line=rgbc(lt['color'], (255, 255, 255)), lw=0.75)
            x0 = lt['x'] + lt['w'] + 4
        w = t['x'] + t['w'] - x0
        fs = pt(t['fs']) * 0.8                       # Arial bold ≈ 1/0.82 de l'Arial Narrow / Avenir Condensed del HTML
        while amplada(t['text'], fs / 0.75, True) / PX > cm(w) and fs > 5: fs -= 0.1
        tb = tbox(sl, cm(x0), cm(t['y']), cm(w), cm(t['h']), 'Titol ' + (t['lt']['text'] if t['lt'] else 'x'), MSO_ANCHOR.MIDDLE)
        r = tb.text_frame.paragraphs[0].add_run(); r.text = t['text']; set_run(r, fs, rgbc(t['color'], (255, 255, 255)), True, TXT_FONT)

    # --- imatges
    for n, im in enumerate(P['imgs'], 1):
        path = ruta_local(im['src'])
        if not os.path.exists(path): avisos.append('imatge no trobada ' + path); continue
        bx, by, bw, bh = im['x'], im['y'], im['w'], im['h']
        ar = im['nw'] / im['nh'] if im['nh'] else 1
        if im['fit'] == 'cover':
            if bw / bh > ar: ww, hh = bw, bw / ar
            else: ww, hh = bh * ar, bh
            f, _ = prepara_imatge(path, cm(ww), cm(hh), tmpdir)
            pic = sl.shapes.add_picture(f, Cm(cm(bx)), Cm(cm(by)), Cm(cm(bw)), Cm(cm(bh)))
            pic.crop_left = pic.crop_right = max(0.0, (ww - bw) / ww / 2); pic.crop_top = pic.crop_bottom = max(0.0, (hh - bh) / hh / 2)
        else:
            if bw / bh > ar: ww, hh = bh * ar, bh
            else: ww, hh = bw, bw / ar
            posx = 0.5; posy = 0.5
            m = re.match(r'([\d.]+)% ([\d.]+)%', im['pos'] or '')
            if m: posx, posy = float(m.group(1)) / 100, float(m.group(2)) / 100
            px_, py_ = bx + (bw - ww) * posx, by + (bh - hh) * posy
            f, _ = prepara_imatge(path, cm(ww), cm(hh), tmpdir)
            pic = sl.shapes.add_picture(f, Cm(cm(px_)), Cm(cm(py_)), Cm(cm(ww)), Cm(cm(hh)))
        pic.name = 'Foto %d' % n
    # --- vores de blocs (mode blocs)
    for n, b in enumerate(P['bloques'], 1):
        rect(sl, cm(b['x']), cm(b['y']), cm(b['w']), cm(b['h']), None, rgbc(b['line']), max(0.5, b['lw'] * 0.75), 'Vora %d' % n)
    # --- referències grans damunt de les fotos
    for r in P['bigs']:
        if not r['round']:
            etiqueta(sl, cm(r['x']), cm(r['y']), cm(r['w']), cm(r['h']), r['text'], pt(r['fs']),
                     rgbc(r['color']), 'Ref ' + r['text'], bold=True, fill=rgbc(r['bg'], (255, 255, 255)),
                     line=rgbc(r['line']), lw=max(0.5, r['lw'] * 0.75))
            continue
        # cercle: l'el·lipse només dona al text el quadrat inscrit i Slides trenca «10» en dues línies -> el número és un
        # quadre de text a part (el doble d'ample, centrat) agrupat amb el cercle
        d = max(r['w'], r['h'], 26.0)                     # diàmetre mínim 26 px (0,69 cm)
        cx, cy = r['x'] + r['w'] / 2, r['y'] + r['h'] / 2
        g = sl.shapes.add_group_shape(); g.name = 'Ref ' + r['text']
        o = g.shapes.add_shape(MSO_SHAPE.OVAL, Cm(cm(cx - d / 2)), Cm(cm(cy - d / 2)), Cm(cm(d)), Cm(cm(d)))
        st = o._element.find(qn('p:style'))
        if st is not None: o._element.remove(st)
        o.fill.solid(); o.fill.fore_color.rgb = rgbc(r['bg'], (255, 255, 255))
        o.line.color.rgb = rgbc(r['line']); o.line.width = Pt(max(0.5, r['lw'] * 0.75)); o.name = 'Cercle ' + r['text']
        t = g.shapes.add_textbox(Cm(cm(cx - d)), Cm(cm(cy - d / 2)), Cm(cm(2 * d)), Cm(cm(d))); t.name = 'Numero ' + r['text']
        tf = t.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        bp = tf._txBody.find(qn('a:bodyPr'))
        for c in list(bp): bp.remove(c)
        etree.SubElement(bp, qn('a:noAutofit'))
        p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
        rn = p0.add_run(); rn.text = r['text']
        fs_r = pt(r['fs'])
        if len(r['text']) > 1: fs_r = min(fs_r, pt(d) * 0.50)          # «10»/«11» amb marge dins del cercle
        set_run(rn, fs_r, rgbc(r['color']), True, TXT_FONT)
    # --- columnes de text
    escala = tria_escala(P['pars'])
    grups = {}
    for p in P['pars']:
        # una columna del HTML pot repartir-se en diverses columnes de la pàgina (mode fluid): agrupem per columna d'origen i x
        grups.setdefault((p['col'], round((p['x']) / 4)), []).append(p)
    nb = 0
    for k, ps in sorted(grups.items(), key=lambda kv: (round(kv[1][0]['x'] / 20), kv[1][0]['y'])):
        ps.sort(key=lambda p: p['y'])
        x = min(p['x'] for p in ps); y = ps[0]['y']; w = max(p['x'] + p['w'] for p in ps) - x
        h = ps[-1]['y'] + ps[-1]['h'] - y
        nb += 1
        tb = tbox(sl, cm(x), cm(y), cm(w), cm(h) + 0.05, 'Text %d' % nb)
        e_g, no_cap = escala_grup(P, ps, escala, x, w, y)
        if e_g < escala - 1e-6: avisos.append('text %d: lletra %.2f -> %.2f per cabre (%.1f pt)' % (nb, escala, e_g, pt(min(p['fs'] for p in ps)) * e_g))
        if no_cap: avisos.append('text %d NO cap ni a la mida mínima' % nb)
        omple_pars(tb, ps, e_g)
    for i, lt in enumerate(P['lts']):
        etiqueta(sl, cm(lt['x']), cm(lt['y']), cm(lt['w']), cm(lt['h']), lt['text'], pt(lt['fs']) * 0.9, rgbc(lt['color']), 'Quadret idioma ' + lt['text'],
                 bold=True, line=rgbc(lt['color']), lw=0.75)

    prs.part._element.find(qn('p:sldSz')).attrib.pop('type', None)
    prs.save(sortida)
    return escala, avisos


def comprova_text(P, ruta):
    """Compara el text de la pàgina HTML amb el del .pptx (paraules): retorna les que falten."""
    prs = Presentation(ruta); s = []
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame: s.append(sh.text_frame.text)
    pptx = re.sub(r'[①-⑳ ]', ' ', ' '.join(s))
    from collections import Counter
    letres = lambda t: Counter(re.findall(r'[^\W\d_]', t.lower()))
    h = re.sub(r'(^|\s)[EFD](?=\s|$)', ' ', P['text'])      # quadrets d'idioma (E/F/D) són formes a part
    falten = letres(h) - letres(pptx)
    return dict(falten)


def previsualitza(pptx_path, png_path, tmp):
    pdfdir = tempfile.mkdtemp(prefix='pptx-prev-')
    try:
        subprocess.run(['/usr/local/bin/soffice', '--headless', '--convert-to', 'pdf', '--outdir', pdfdir, pptx_path],
                       capture_output=True, timeout=180)
        pdf = os.path.join(pdfdir, os.path.splitext(os.path.basename(pptx_path))[0] + '.pdf')
        base = png_path[:-4]
        subprocess.run(['pdftoppm', '-r', '70', '-png', '-singlefile', pdf, base], check=True)
    finally:
        shutil.rmtree(pdfdir, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--grup', type=int, required=True)
    ap.add_argument('--pagina', type=int, action='append', help='número de pàgina renderitzada (repetible); defecte: totes les d\'informació')
    ap.add_argument('--html', help='defecte: sortida/gN.html')
    ap.add_argument('--sortida', default=os.path.join(ROOT, 'slides'))
    ap.add_argument('--exclou', type=int, action='append', default=None,
                    help='pàgines a saltar (defecte: les que ja són a slides-pagines.yaml per al grup)')
    ap.add_argument('--previsualitza', action='store_true')
    a = ap.parse_args()
    html_path = a.html or os.path.join(ROOT, 'sortida', 'g%d.html' % a.grup)
    exclou = a.exclou
    if exclou is None:
        exclou = []
        try:
            import yaml
            y = yaml.safe_load(open(os.path.join(ROOT, 'slides-pagines.yaml'), encoding='utf8')) or {}
            exclou = [int(k) for k in (y.get(a.grup) or y.get(str(a.grup)) or {}).keys()]
        except Exception: pass
    pagines = extreu(html_path)
    os.makedirs(a.sortida, exist_ok=True)
    tmpdir = tempfile.mkdtemp(prefix='pptx-img-')
    try:
        for P in pagines:
            if a.pagina and P['n'] not in a.pagina: continue
            if not a.pagina and P['n'] in exclou: continue
            nom = 'G%d-p%02d' % (a.grup, P['n'])
            ruta = os.path.join(a.sortida, nom + '.pptx')
            escala, avisos = construeix(P, ruta, tmpdir)
            falten = comprova_text(P, ruta)
            mb = os.path.getsize(ruta) / 1e6
            print('%s  %s  %.2f MB  escala lletra %.2f  par=%d fotos=%d%s%s' % (
                nom, P['classes'].strip(), mb, escala, len(P['pars']), len(P['imgs']),
                '  AVIS ' + '; '.join(avisos) if avisos else '', '  TEXT-QUE-FALTA ' + str(falten) if falten else ''))
            if mb >= 8: print('  !! >= 8 MB')
            if a.previsualitza: previsualitza(ruta, os.path.join(a.sortida, nom + '-previsualitzacio.png'), tmpdir)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == '__main__':
    main()
